#!/usr/bin/env python
"""
VM Field Verification Script
อ่านและตรวจสอบ field Owner / Department / Environment / Project / Server
จากไฟล์ Excel ที่ระบุใน config.json โดย 1 แถวแทน 1 VM

ใช้: python 01_vm_field_verify.py [path_to_xlsx]
Output:
  - .output/01_vm_field_audit.csv
  - .output/01_vm_all_data.csv            (raw snapshot ทุก VM ที่ scan ได้ ไม่กรอง)
  - .output/01_vm_owner_email_list.csv    (owner email ทุกคนแบบ unique + จำนวน VM ที่เป็น owner)
"""

import sys
import re
import csv
from collections import defaultdict

import openpyxl
from vm_config import add_pending_vms, load_survey_config

# =====================================================================
# Survey settings are kept in the local, ignored config.json.
# =====================================================================

OUTPUT_DIR = ".output"

SURVEY_CONFIG = load_survey_config()
INPUT_FILE_DEFAULT = SURVEY_CONFIG["SURVEY_FILE"]
DEPARTMENT_SHEETS = SURVEY_CONFIG["DEPARTMENT_SHEETS"]

# header อยู่แถว 4, EXAMPLE แถว 3, ข้อมูลจริงเริ่มแถว 5
HEADER_ROW = 4
DATA_START_ROW = 5

# Optional sanity check for the total number of VMs in the survey.
EXPECTED_TOTAL_VM = SURVEY_CONFIG.get("EXPECTED_VM_COUNT")

# mapping คอลัมน์ (1-indexed ตาม openpyxl)
COL_VM_NAME = 1  # A — VM Name
COL_RG_NAME = 2  # B — Resource Group
COL_OWNER = 3  # C — owner
COL_DEPARTMENT = 4  # D — department
COL_ENVIRONMENT = 5  # E — environment
COL_PROJECT = 6  # F — project
COL_SERVER = 7  # G — server (optional)
COL_DATA_CONFIRMED = 8  # H — Data Confirmed? (Confirmed / Pending)
COL_ACTIVE_VM = 9  # I — Active VM? (Active / Inactive / Unknown)

# ค่าที่ถือว่า "Inactive" ในคอลัมน์ Active VM? — เวลาช่องข้อมูลว่างเปล่า
# ให้เช็คด้วยว่า VM นั้นถูกมาร์คไว้เป็น Inactive หรือไม่ ถ้าใช่ให้ระบุไว้ใน Issue
INACTIVE_VALUE = "Inactive"

OWNER_DOMAIN_WHITELIST = {value.lower() for value in SURVEY_CONFIG["OWNER_DOMAIN_WHITELIST"]}
DEPARTMENT_MASTER_LIST = set(SURVEY_CONFIG["DEPARTMENT_MASTER_LIST"])
ENVIRONMENT_MASTER_LIST = set(SURVEY_CONFIG["ENVIRONMENT_MASTER_LIST"])

# regex email ทั่วไป local@domain.tld
EMAIL_REGEX = re.compile(r"^[^\s@,]+@[^\s@,]+\.[^\s@,]+$")

SEVERITY_ORDER = {"High": 0, "Medium": 1, "Low": 2}


# =====================================================================
# HELPERS
# =====================================================================

def is_data_row(ws, row_idx):
    """เอาเฉพาะแถวที่คอลัมน์ A (VM Name) ไม่ว่าง — 1 VM Name ที่ไม่ว่าง = 1 record"""
    val = ws.cell(row=row_idx, column=COL_VM_NAME).value
    return val is not None and str(val).strip() != ""


def split_multi_value(raw):
    """แยกค่าที่คั่นด้วย comma (owner/project อาจมีหลายค่าใน 1 cell)"""
    if raw is None:
        return []
    text = str(raw)
    return text.split(",")


def annotate_if_inactive(issues, active_vm_value):
    """
    ถ้า issue เป็น "...is empty" และ VM ถูกมาร์คไว้เป็น Inactive
    (คอลัมน์ Active VM? = "Inactive") ให้เติม note บอกไว้ใน issue text
    ไม่เปลี่ยน severity — แค่ระบุข้อมูลเพิ่มให้คนอ่านตัดสินใจเอง
    """
    is_inactive = (
        active_vm_value is not None and str(active_vm_value).strip() == INACTIVE_VALUE
    )
    if not is_inactive:
        return issues

    annotated = []
    for issue_text, severity in issues:
        if issue_text.endswith("is empty"):
            issue_text = f"{issue_text} (VM marked Inactive)"
        annotated.append((issue_text, severity))
    return annotated


def has_bad_comma_spacing(raw):
    """เช็คว่ามี space หลัง comma หรือ space หัว-ท้ายหรือไม่ (raw string ทั้งก้อน)"""
    if raw is None:
        return False
    text = str(raw)
    if text != text.strip():
        return True
    if ", " in text:
        return True
    return False


# =====================================================================
# FIELD CHECKERS — แต่ละฟังก์ชัน return list of (issue_text, severity)
# เพิ่ม field ใหม่ในอนาคต: เพิ่มฟังก์ชัน checker + map เข้า SIMPLE_FIELD_CHECKERS ด้านล่าง
# =====================================================================

def check_owner(raw_value):
    """Owner checks"""
    issues = []

    # Owner ว่างเปล่า -> High
    if raw_value is None or str(raw_value).strip() == "":
        issues.append(("Owner is empty", "High"))
        return issues

    # space หลัง comma หรือ space หัว-ท้าย -> Medium
    if has_bad_comma_spacing(raw_value):
        issues.append(("Owner has extra space after comma or leading/trailing space", "Medium"))

    # comma ลงท้าย / นำหน้า / ซ้อนกัน (มี entry ว่างระหว่าง comma) -> Medium
    if any(p.strip() == "" for p in split_multi_value(raw_value)):
        issues.append(("Owner has empty entry (leading/trailing/double comma)", "Medium"))

    emails = [e.strip() for e in split_multi_value(raw_value) if e.strip() != ""]

    for email in emails:
        # ไม่ตรง regex email ทั่วไป -> High
        if not EMAIL_REGEX.match(email):
            issues.append((f"Owner value '{email}' does not match email format", "High"))
            continue  # ไม่มี domain ให้เช็คต่อถ้า format ผิด

        domain = email.split("@", 1)[1].lower()

        # domain ไม่อยู่ใน whitelist -> High
        if domain not in OWNER_DOMAIN_WHITELIST:
            issues.append((f"Owner email '{email}' domain not in whitelist ({domain})", "High"))

    return issues


def check_free_text(raw_value, field_label, empty_severity="High"):
    """
    Free-text checks (ใช้กับ Project และ Server)
    ว่างเปล่า -> empty_severity (Project = High, Server = Low เพราะ optional)
    มี space หลัง comma หรือ space หัว-ท้าย -> Low
    """
    issues = []

    # ว่างเปล่า -> empty_severity
    if raw_value is None or str(raw_value).strip() == "":
        issues.append((f"{field_label} is empty", empty_severity))
        return issues

    # มี space หลัง comma -> Low
    if has_bad_comma_spacing(raw_value):
        issues.append(
            (f"{field_label} has extra space after comma or leading/trailing space", "Low")
        )

    return issues


def check_project(raw_value):
    return check_free_text(raw_value, "Project")


def check_server(raw_value):
    # server เป็น optional — ว่างเปล่าถือเป็น Low
    return check_free_text(raw_value, "Server", empty_severity="Low")


def check_against_master_list(raw_value, master_list, field_label):
    """
    Department / Environment checks
    ว่างเปล่า -> High, ไม่ตรง exact match -> Medium
    ไม่มี fuzzy-match / cross-reference ใดๆ
    """
    issues = []

    if raw_value is None or str(raw_value).strip() == "":
        issues.append((f"{field_label} is empty", "High"))
        return issues

    value = str(raw_value).strip()
    if value not in master_list:
        issues.append((f"{field_label} value '{value}' not an exact match in master list", "Medium"))

    return issues


def check_department(raw_value):
    return check_against_master_list(raw_value, DEPARTMENT_MASTER_LIST, "Department")


def check_environment(raw_value):
    return check_against_master_list(raw_value, ENVIRONMENT_MASTER_LIST, "Environment")


# Field registry — เพิ่ม field ใหม่ในอนาคต: เพิ่ม checker ด้านบน แล้ว map ที่นี่
# (owner ไม่ได้อยู่ใน registry นี้ เพราะรวม field ที่มี logic ต่างกันเยอะกว่า — จัดการแยกใน main loop)
SIMPLE_FIELD_CHECKERS = {
    "Project": (COL_PROJECT, check_project),
    "Server": (COL_SERVER, check_server),
    "Department": (COL_DEPARTMENT, check_department),
    "Environment": (COL_ENVIRONMENT, check_environment),
}


# =====================================================================
# MAIN
# =====================================================================

def main():
    input_path = sys.argv[1] if len(sys.argv) > 1 else INPUT_FILE_DEFAULT

    wb = openpyxl.load_workbook(input_path, data_only=True)

    audit_rows = []  # rows for .output/01_vm_field_audit.csv
    all_vm_rows = []  # rows for .output/01_vm_all_data.csv — every VM scanned, no filtering
    pending_vm_names = []
    # เก็บ owner email ทุกตัวที่เจอ (ไม่ว่า format จะถูกหรือผิด) ไว้ทำ .output/01_vm_owner_email_list.csv
    # key = email lowercase+strip, value = set of VM names ที่เจอ email นี้
    all_owner_emails = defaultdict(set)

    total_vm_scanned = 0
    # เก็บจำนวน VM ที่ scan ได้แยกต่อ sheet เพื่อไล่หา sheet ที่ผิดได้ถ้ารวมไม่ตรง EXPECTED_TOTAL_VM
    vm_count_by_sheet = {}
    issue_count_by_field = defaultdict(int)
    issue_count_by_severity = defaultdict(int)

    for sheet_name in DEPARTMENT_SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"[WARN] Sheet not found, skipped: {sheet_name}")
            vm_count_by_sheet[sheet_name] = 0
            continue
        ws = wb[sheet_name]
        sheet_vm_count = 0

        for row_idx in range(DATA_START_ROW, ws.max_row + 1):
            if not is_data_row(ws, row_idx):
                continue

            total_vm_scanned += 1
            sheet_vm_count += 1
            vm_name = str(ws.cell(row=row_idx, column=COL_VM_NAME).value).strip()
            rg_name = ws.cell(row=row_idx, column=COL_RG_NAME).value
            data_confirmed_value = ws.cell(row=row_idx, column=COL_DATA_CONFIRMED).value
            active_vm_value = ws.cell(row=row_idx, column=COL_ACTIVE_VM).value
            if data_confirmed_value is not None and str(data_confirmed_value).strip().casefold() == "pending":
                pending_vm_names.append(vm_name)

            # ---- Owner ----
            owner_raw = ws.cell(row=row_idx, column=COL_OWNER).value
            owner_issues = check_owner(owner_raw)
            owner_issues = annotate_if_inactive(owner_issues, active_vm_value)
            for issue_text, severity in owner_issues:
                audit_rows.append({
                    "VM Name": vm_name,
                    "Resource Group": "" if rg_name is None else str(rg_name),
                    "Department (Sheet)": sheet_name,
                    "Field": "Owner",
                    "Value": "" if owner_raw is None else str(owner_raw),
                    "Issue": issue_text,
                    "Severity": severity,
                })
                issue_count_by_field["Owner"] += 1
                issue_count_by_severity[severity] += 1
            # เก็บทุก entry owner email (ไม่ว่า format จะถูกหรือผิด) สำหรับ owner email list รวม
            if owner_raw:
                for email in [e.strip() for e in split_multi_value(owner_raw) if e.strip()]:
                    all_owner_emails[email.lower()].add(vm_name)

            # ---- Project / Server / Department / Environment ----
            for field_label, (col, checker) in SIMPLE_FIELD_CHECKERS.items():
                raw_val = ws.cell(row=row_idx, column=col).value
                field_issues = checker(raw_val)
                field_issues = annotate_if_inactive(field_issues, active_vm_value)
                for issue_text, severity in field_issues:
                    audit_rows.append({
                        "VM Name": vm_name,
                        "Resource Group": "" if rg_name is None else str(rg_name),
                        "Department (Sheet)": sheet_name,
                        "Field": field_label,
                        "Value": "" if raw_val is None else str(raw_val),
                        "Issue": issue_text,
                        "Severity": severity,
                    })
                    issue_count_by_field[field_label] += 1
                    issue_count_by_severity[severity] += 1

            # ---- Full VM data snapshot — every VM scanned, raw values, no filtering ----
            department_raw = ws.cell(row=row_idx, column=COL_DEPARTMENT).value
            environment_raw = ws.cell(row=row_idx, column=COL_ENVIRONMENT).value
            project_raw = ws.cell(row=row_idx, column=COL_PROJECT).value
            server_raw = ws.cell(row=row_idx, column=COL_SERVER).value
            all_vm_rows.append({
                "VM Name": vm_name,
                "Resource Group": "" if rg_name is None else str(rg_name),
                "Owner (Email)": "" if owner_raw is None else str(owner_raw),
                "Department": "" if department_raw is None else str(department_raw),
                "Environment": "" if environment_raw is None else str(environment_raw),
                "Project": "" if project_raw is None else str(project_raw),
                "Server": "" if server_raw is None else str(server_raw),
                "Data Confirmed?": "" if data_confirmed_value is None else str(data_confirmed_value),
                "Active VM?": "" if active_vm_value is None else str(active_vm_value),
            })

        # บันทึกจำนวน VM ที่ scan ได้ของ sheet นี้
        vm_count_by_sheet[sheet_name] = sheet_vm_count

    # sort audit rows: by Field then Severity (High->Medium->Low) then VM Name
    audit_rows.sort(key=lambda r: (r["Field"], SEVERITY_ORDER.get(r["Severity"], 9), r["VM Name"]))

    empty_on_inactive_count = sum(
        1 for r in audit_rows if "(VM marked Inactive)" in r["Issue"]
    )


    # ---- Output 1: .output/01_vm_field_audit.csv ----
    audit_csv_path = ".output/01_vm_field_audit.csv"
    with open(audit_csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "VM Name", "Resource Group", "Department (Sheet)",
                "Field", "Value", "Issue", "Severity",
            ],
        )
        writer.writeheader()
        writer.writerows(audit_rows)

    # ---- Output 2: .output/01_vm_all_data.csv (full snapshot ทุก VM ไม่กรอง) ----
    all_vm_csv_path = ".output/01_vm_all_data.csv"
    with open(all_vm_csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "VM Name",
                "Resource Group",
                "Owner (Email)",
                "Department",
                "Environment",
                "Project",
                "Server",
                "Data Confirmed?",
                "Active VM?",
            ],
        )
        writer.writeheader()
        writer.writerows(all_vm_rows)

    added_skip_vm, total_skip_vm = add_pending_vms(pending_vm_names)
    print(f"Pending VM: {len(pending_vm_names)}; added to config.json SKIP_VM: {added_skip_vm}; total: {total_skip_vm}")

    # ---- Output 3: .output/01_vm_owner_email_list.csv (owner ทุกคน unique + count) ----
    owner_email_rows = []
    for email in sorted(all_owner_emails.keys()):
        vm_list = sorted(all_owner_emails[email])
        owner_email_rows.append({
            "email": email,
            "จำนวน VM ที่เป็น owner": len(vm_list),
            "VM list": ", ".join(vm_list),
        })
    # เรียงคนที่เป็น owner หลาย VM ที่สุดขึ้นก่อน
    owner_email_rows.sort(key=lambda r: -r["จำนวน VM ที่เป็น owner"])

    owner_email_csv_path = ".output/01_vm_owner_email_list.csv"
    with open(owner_email_csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f, fieldnames=["email", "จำนวน VM ที่เป็น owner", "VM list"]
        )
        writer.writeheader()
        writer.writerows(owner_email_rows)

    # ---- Console summary ----
    print("=" * 60)
    print("VM Field Verification Summary")
    print("=" * 60)
    print(f"VM scanned (total across {len(DEPARTMENT_SHEETS)} department sheets): {total_vm_scanned}")
    for sheet_name in DEPARTMENT_SHEETS:
        print(f"  - {sheet_name:32s}: {vm_count_by_sheet.get(sheet_name, 0)}")
    print()

    # ---- optional survey size check ----
    if EXPECTED_TOTAL_VM is not None:
        if total_vm_scanned == EXPECTED_TOTAL_VM:
            print(f"VM count check: {total_vm_scanned}/{EXPECTED_TOTAL_VM} OK")
        else:
            diff = total_vm_scanned - EXPECTED_TOTAL_VM
            diff_label = f"+{diff}" if diff > 0 else str(diff)
            print("!" * 60)
            print(f"[WARNING] VM count MISMATCH: scanned {total_vm_scanned}, expected {EXPECTED_TOTAL_VM} ({diff_label})")
            print("Check per-sheet breakdown above for which sheet(s) look off")
            print("(e.g. a VM Name cell is blank, or data doesn't start at row 5)")
            print("!" * 60)
    print()
    print("Issues by field:")
    for field in ["Owner", "Project", "Server", "Department", "Environment"]:
        print(f"  {field:12s}: {issue_count_by_field.get(field, 0)}")
    print()
    print("Issues by severity:")
    for sev in ["High", "Medium", "Low"]:
        print(f"  {sev:8s}: {issue_count_by_severity.get(sev, 0)}")
    print()
    print(f"Total issues: {sum(issue_count_by_severity.values())}")
    print(f"  of which 'empty field' on VM marked Inactive: {empty_on_inactive_count}")
    print(f"Unique owner emails found (all, incl. flagged): {len(all_owner_emails)}")
    print()
    print("Output written:")
    print(f"  - {audit_csv_path}")
    print(f"  - {all_vm_csv_path}")
    print(f"  - {owner_email_csv_path}")


if __name__ == "__main__":
    main()
